#!/usr/bin/env python3
"""Empirically determine how ``pptx_from_template.py`` maps layouts.

Answers one question by test, not assumption: does ``--add-slides`` accept a
template's own layout *names*, or only the built-in keys?

The probe reports success only when the OUTPUT deck's slide count goes up by
exactly one. Exit code 0 alone is not success — the ``--values``/``--add-slides``
trap means a spec passed to the wrong flag exits 0 and appends nothing.
"""

import argparse
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation


def _try_build(template, from_template, python, runner, layout):
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "out.pptx"
        spec = Path(tmp) / "spec.json"
        spec.write_text(
            json.dumps({"slides": [{"layout": layout, "title": "Layout test"}]})
        )
        r = runner(
            [python, from_template, str(template), str(out), "--add-slides", str(spec)],
            capture_output=True,
            text=True,
        )
        if r.returncode != 0:
            return False
        try:
            n_out = len(Presentation(str(out)).slides)
            n_tpl = len(Presentation(str(template)).slides)
        except Exception:
            return False
        return n_out == n_tpl + 1


def probe(template, scripts_dir, python=None, runner=None):
    """Return ``{"mode", "layouts_available", "tested_layout"}``.

    ``mode`` is one of ``template-names``, ``builtin-keys``, or ``failed``.
    """
    if python is None:
        python = sys.executable
    if runner is None:
        runner = subprocess.run

    scripts_dir = Path(scripts_dir)
    template = Path(template)
    read = str(scripts_dir / "pptx_read.py")
    from_template = str(scripts_dir / "pptx_from_template.py")

    r = runner(
        [python, read, str(template), "--outline"], capture_output=True, text=True
    )
    if r.returncode != 0:
        return {"mode": "failed", "layouts_available": [], "tested_layout": ""}
    try:
        data = json.loads(r.stdout)
    except Exception:
        return {"mode": "failed", "layouts_available": [], "tested_layout": ""}
    layouts = list(data.get("layouts_available", []))

    if not layouts:
        return {"mode": "failed", "layouts_available": layouts, "tested_layout": ""}

    tested = next((l for l in layouts if "blank" not in l.lower()), layouts[0])

    if _try_build(template, from_template, python, runner, tested):
        return {
            "mode": "template-names",
            "layouts_available": layouts,
            "tested_layout": tested,
        }

    if tested != "title_content" and _try_build(
        template, from_template, python, runner, "title_content"
    ):
        return {
            "mode": "builtin-keys",
            "layouts_available": layouts,
            "tested_layout": "title_content",
        }

    return {"mode": "failed", "layouts_available": layouts, "tested_layout": tested}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", help="path to the template .pptx")
    parser.add_argument("--scripts-dir", required=True, help="dir containing pptx_*.py")
    parser.add_argument(
        "--out", required=True, help="cache file to write {mode, layouts_available}"
    )
    args = parser.parse_args(argv)

    result = probe(args.template, args.scripts_dir)
    cache = {"mode": result["mode"], "layouts_available": result["layouts_available"]}
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(cache, indent=2) + "\n")
    print(json.dumps(result, indent=2))
    return 1 if result["mode"] == "failed" else 0


if __name__ == "__main__":
    sys.exit(main())
