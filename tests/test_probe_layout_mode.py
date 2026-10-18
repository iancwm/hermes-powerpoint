import json
import pathlib
import shutil
import subprocess

import pytest
from pptx import Presentation

from provision.probe_layout_mode import main, probe

HOME = pathlib.Path.home()


def _scripts_dir():
    skills = HOME / ".hermes" / "skills"
    if not skills.exists():
        return None
    matches = list(skills.rglob("pptx_from_template.py"))
    return matches[0].parent if matches else None


@pytest.fixture
def scripts_dir():
    d = _scripts_dir()
    if d is None:
        pytest.skip("Hermes powerpoint skill scripts not installed")
    return d


@pytest.fixture
def make_template():
    def _make(tmp_path, rename=None):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        if rename is not None:
            prs.slide_layouts[0].name = rename
        p = tmp_path / "template.pptx"
        prs.save(str(p))
        return p

    return _make


def test_default_template_uses_template_names(tmp_path, scripts_dir, make_template):
    template = make_template(tmp_path)
    result = probe(template, scripts_dir)
    assert result["mode"] == "template-names"
    assert "Title and Content" in result["layouts_available"]


def test_layout_name_with_quotes_and_ampersand(tmp_path, scripts_dir, make_template):
    template = make_template(tmp_path, rename='Title "Quoted" & Content')
    result = probe(template, scripts_dir)
    assert result["tested_layout"] == 'Title "Quoted" & Content'
    assert result["mode"] == "template-names"


def test_exit_zero_without_new_slide_is_not_success(
    tmp_path, scripts_dir, make_template, monkeypatch
):
    template = make_template(tmp_path)
    real_run = subprocess.run

    def fake_run(argv, *a, **kw):
        if "--add-slides" in argv:
            shutil.copy(argv[2], argv[3])
            return subprocess.CompletedProcess(argv, 0, stdout="", stderr="")
        return real_run(argv, *a, **kw)

    monkeypatch.setattr("provision.probe_layout_mode.subprocess.run", fake_run)
    result = probe(template, scripts_dir)
    assert result["mode"] == "failed"


def test_cli_writes_cache_file(tmp_path, scripts_dir, make_template):
    template = make_template(tmp_path)
    out = tmp_path / "brand-layouts.json"
    rc = main([str(template), "--scripts-dir", str(scripts_dir), "--out", str(out)])
    assert rc == 0
    data = json.loads(out.read_text())
    assert set(data.keys()) == {"mode", "layouts_available"}
    assert "slides" not in data
